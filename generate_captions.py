"""Caption DOCX generator for Korean variety-style subtitles.

This script stitches together analyzed subtitle candidates, optional program
lexicon hints, and caption theme presets to emit a Word document (.docx) that
follows the style rules outlined in README.md. It focuses on deterministic
styling so PD guide captions and buildup sequences remain consistent.
"""
from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


DEFAULT_STYLES = {
    "고딕": {"font_name": "Malgun Gothic", "size": 20, "align": WD_ALIGN_PARAGRAPH.LEFT},
    "고딕 중앙": {"font_name": "Malgun Gothic", "size": 24, "align": WD_ALIGN_PARAGRAPH.CENTER},
    "중앙대판": {"font_name": "Malgun Gothic", "size": 32, "align": WD_ALIGN_PARAGRAPH.CENTER},
    "굴림": {"font_name": "Gulim", "size": 18, "align": WD_ALIGN_PARAGRAPH.LEFT},
}

SIZE_MAP = {
    "small": 18,
    "medium": 24,
    "large": 32,
}


@dataclass
class CaptionClip:
    text: str
    subtitle_type: str
    font: Optional[str] = None
    align: Optional[str] = None
    size: Optional[str] = None
    sequence_tag: Optional[str] = None
    tone_pack: Optional[str] = None
    notes: Optional[str] = None
    keep_original: bool = False
    hook: bool = False
    cta: bool = False
    theme_applied: Optional[str] = None


@dataclass
class CaptionTheme:
    style_overrides: Dict[str, Dict[str, str]] = field(default_factory=dict)
    style_id: Optional[str] = None

    @classmethod
    def from_file(cls, path: Optional[Path]) -> "CaptionTheme":
        if not path:
            return cls()
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        return cls(style_overrides=data.get("styles", {}), style_id=data.get("style_id"))

    def resolve_style(self, clip: CaptionClip) -> Dict[str, str]:
        preferred_font = clip.font or self.style_overrides.get(clip.subtitle_type, {}).get("font")
        style_name = preferred_font or self.subtitle_type_default(clip.subtitle_type)
        return {
            "style_name": style_name,
            "align": clip.align or self.style_overrides.get(style_name, {}).get("align"),
            "size": clip.size or self.style_overrides.get(style_name, {}).get("size"),
        }

    @staticmethod
    def subtitle_type_default(subtitle_type: str) -> str:
        mapping = {
            "말": "고딕",
            "상황": "굴림",
            "재미": "고딕 중앙",
        }
        return mapping.get(subtitle_type, "고딕")


@dataclass
class ProgramLexicon:
    names: List[str] = field(default_factory=list)
    terms: List[str] = field(default_factory=list)

    @classmethod
    def from_file(cls, path: Optional[Path]) -> "ProgramLexicon":
        if not path:
            return cls()
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        names = []
        for item in data.get("names", []):
            if isinstance(item, dict):
                names.append(item.get("label") or item.get("name"))
                names.extend(item.get("alias", []))
            elif isinstance(item, str):
                names.append(item)
        terms = data.get("terms", [])
        return cls(names=[n for n in names if n], terms=terms)

    def annotate(self, clip: CaptionClip) -> Optional[str]:
        markers = []
        for name in self.names:
            if name in clip.text:
                markers.append(f"name:{name}")
        for term in self.terms:
            if term in clip.text:
                markers.append(f"term:{term}")
        if markers:
            return (clip.notes or "") + f" | lexicon={','.join(markers)}" if clip.notes else f"lexicon={','.join(markers)}"
        return clip.notes


def load_clips(path: Path) -> List[CaptionClip]:
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    clips = []
    for raw in data.get("clips", []):
        clips.append(
            CaptionClip(
                text=raw["text"],
                subtitle_type=raw.get("subtitle_type", "말"),
                font=raw.get("font"),
                align=raw.get("align"),
                size=raw.get("size"),
                sequence_tag=raw.get("sequence_tag"),
                tone_pack=raw.get("tone_pack"),
                notes=raw.get("notes"),
                keep_original=raw.get("keep_original", False),
                hook=raw.get("hook", False),
                cta=raw.get("cta", False),
                theme_applied=raw.get("theme_applied"),
            )
        )
    return clips


def ensure_styles(doc: Document) -> None:
    for style_name, config in DEFAULT_STYLES.items():
        if style_name in doc.styles:
            continue
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = config["font_name"]
        style.font.size = Pt(config["size"])
        style.paragraph_format.alignment = config["align"]


def apply_alignment(paragraph, align: Optional[str], default_align):
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = default_align


def apply_size(paragraph, size: Optional[str]):
    if not size:
        return
    font_size = SIZE_MAP.get(size, None)
    if font_size:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def add_hidden_metadata(paragraph, clip: CaptionClip):
    meta_parts = [f"type={clip.subtitle_type}"]
    if clip.sequence_tag:
        meta_parts.append(f"sequence={clip.sequence_tag}")
    if clip.hook:
        meta_parts.append("hook=true")
    if clip.cta:
        meta_parts.append("CTA=true")
    if clip.keep_original:
        meta_parts.append("keep_original=true")
    if clip.tone_pack:
        meta_parts.append(f"tone_pack={clip.tone_pack}")
    if clip.theme_applied:
        meta_parts.append(f"theme={clip.theme_applied}")
    if clip.notes:
        meta_parts.append(f"notes={clip.notes}")
    if meta_parts:
        meta_run = paragraph.add_run(f" [meta: {';'.join(meta_parts)}]")
        meta_run.font.hidden = True
        meta_run.font.size = Pt(1)


def write_document(clips: List[CaptionClip], theme: CaptionTheme, lexicon: ProgramLexicon, output_path: Path):
    doc = Document()
    ensure_styles(doc)

    for clip in clips:
        clip.notes = lexicon.annotate(clip)
        style_info = theme.resolve_style(clip)
        style_name = style_info["style_name"]
        style = doc.styles.get(style_name)
        paragraph = doc.add_paragraph(style=style)
        paragraph.add_run(clip.text)
        apply_alignment(paragraph, style_info.get("align"), style.paragraph_format.alignment)
        apply_size(paragraph, style_info.get("size"))
        add_hidden_metadata(paragraph, clip)

    doc.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate variety-style caption docx from analyzed clips.")
    parser.add_argument("--analysis", type=Path, required=True, help="Path to analyzed clips JSON")
    parser.add_argument("--lexicon", type=Path, help="Optional program lexicon JSON")
    parser.add_argument("--theme", type=Path, help="Optional caption theme JSON")
    parser.add_argument("--output", type=Path, default=Path("captions.docx"), help="Output .docx path")
    return parser.parse_args()


def main():
    args = parse_args()
    clips = load_clips(args.analysis)
    theme = CaptionTheme.from_file(args.theme)
    lexicon = ProgramLexicon.from_file(args.lexicon)
    write_document(clips, theme, lexicon, args.output)
    print(f"Wrote {len(clips)} captions to {args.output}")


if __name__ == "__main__":
    main()
